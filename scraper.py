"""
Org Finder - CRAWL stage. Discovers and fetches organisation websites and
upserts their raw text + contact details into a shared SQLite database
(orgs.db). Also does a coarse relevance scoring (criteria.json) just to decide
what's worth snowballing from - this is NOT the final ranking.

For picking parameters and getting a ranked, filtered list of leads, run the
QUERY stage instead: `streamlit run app.py` (reads orgs.db, no re-crawling,
instant). Run this script again any time you want to refresh/expand the
database (it's additive - existing organisations are updated, not lost).

Discovery:
  1. seeds.csv            - organisations you already know of
  2. snowballing          - organisations linked from seeds' partner/network pages
  3. web search (optional)- set SERPER_API_KEY (free: 2,500 one-time queries, https://serper.dev),
                             then pass --search. Query phrases come from params_library.json's
                             per-tag 'search_queries' lists, not a hardcoded list.

Anything already in orgs.db is never re-crawled (pass --retry-errors to retry past failures only).

Usage:
  python scraper.py                      # seeds + 1 level of snowballing -> orgs.db
  python scraper.py --depth 0            # seeds only
  python scraper.py --depth 2 --max-orgs 500
  python scraper.py --search             # also discover via Serper (needs SERPER_API_KEY)
"""

import argparse
import csv
import json
import os
import re
import sys
import threading
import time
import warnings
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from datetime import datetime, timezone
from urllib.parse import urljoin, urlparse
from urllib.robotparser import RobotFileParser

import db as orgdb  # noqa: E402

warnings.filterwarnings("ignore")
import requests  # noqa: E402
from bs4 import BeautifulSoup  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/126.0 Safari/537.36 OrgFinder/1.0")
HEADERS = {"User-Agent": UA, "Accept-Language": "en-IN,en;q=0.9"}
TIMEOUT = 20

# Internal pages worth reading (matched against URL path and link text)
RELEVANT = re.compile(
    r"about|who-?we-?are|what-?we-?do|our-?work|programm?e|initiative|project|approach|"
    r"model|impact|reach|where-?we-?work|fln|literacy|numeracy|foundational|teacher|"
    r"partner|network|collaborat|school|learning", re.I)
PARTNERISH = re.compile(r"partner|network|collaborat|associat|funder|supporter|donor|alliance|ecosystem", re.I)
SKIP_EXT = re.compile(r"\.(pdf|jpe?g|png|gif|svg|webp|zip|docx?|xlsx?|pptx?|mp4|mp3)(\?|$)", re.I)

# Domains never treated as candidate organisations
BLOCKED = {
    "facebook.com", "twitter.com", "x.com", "instagram.com", "linkedin.com", "youtube.com", "youtu.be",
    "google.com", "goo.gl", "apple.com", "microsoft.com", "wikipedia.org", "wordpress.com", "wordpress.org",
    "wix.com", "wixsite.com", "medium.com", "amazon.com", "amazon.in", "whatsapp.com", "wa.me", "bit.ly",
    "t.me", "github.com", "issuu.com", "scribd.com", "flickr.com", "vimeo.com", "pinterest.com",
    "thehindu.com", "indianexpress.com", "timesofindia.com", "indiatimes.com", "hindustantimes.com",
    "livemint.com", "economictimes.com", "ndtv.com", "scroll.in", "thewire.in", "theprint.in",
    "thebetterindia.com", "yourstory.com", "forbes.com", "forbesindia.com", "bbc.com", "bbc.co.uk",
    "business-standard.com", "deccanherald.com", "news18.com", "outlookindia.com", "edexlive.com",
    "give.do", "guidestar.org", "ngoportal.org", "idealist.org", "devex.com", "globalgiving.org",
    "gstatic.com", "googleapis.com", "w3.org", "schema.org", "mailchimp.com", "list-manage.com",
    "razorpay.com", "paypal.com", "instamojo.com", "cloudflare.com", "jquery.com", "elementor.com",
    "squarespace.com", "godaddy.com", "hubspot.com", "zoom.us", "forms.gle", "typeform.com",
    "spotify.com", "soundcloud.com", "tiktok.com", "threads.net", "creativecommons.org",
    "googletagmanager.com", "doubleclick.net", "addtoany.com", "sharethis.com", "wpengine.com",
    "unicef.org", "unesco.org", "worldbank.org", "un.org", "who.int", "nic.in", "gov.in",
}
MULTI_TLD = ("org.in", "co.in", "ac.in", "net.in", "gov.in", "nic.in", "edu.in", "res.in",
             "org.uk", "co.uk", "ac.uk", "com.au", "org.au", "co.za", "org.za")

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?:\+91[\s-]?|0)?[6-9]\d{4}[\s-]?\d{5}|\+91[\s-]?\d{2,4}[\s-]?\d{6,8}")


def reg_domain(url_or_host: str) -> str:
    host = urlparse(url_or_host).netloc if "//" in url_or_host else url_or_host
    host = host.lower().split(":")[0]
    if host.startswith("www."):
        host = host[4:]
    parts = host.split(".")
    if len(parts) >= 3 and ".".join(parts[-2:]) in MULTI_TLD:
        return ".".join(parts[-3:])
    return ".".join(parts[-2:])


def is_blocked(domain: str) -> bool:
    if domain in BLOCKED:
        return True
    return any(domain.endswith("." + b) for b in BLOCKED) or domain.endswith((".gov", ".gov.in", ".nic.in"))


# ---------------------------------------------------------------- fetching

_session_local = threading.local()


def session() -> requests.Session:
    if not hasattr(_session_local, "s"):
        s = requests.Session()
        s.headers.update(HEADERS)
        _session_local.s = s
    return _session_local.s


def fetch(url: str):
    """Return (final_url, html) or raise."""
    r = session().get(url, timeout=TIMEOUT, allow_redirects=True)
    r.raise_for_status()
    ctype = r.headers.get("content-type", "")
    if "html" not in ctype and "xml" not in ctype and ctype:
        raise ValueError(f"not html ({ctype})")
    enc = r.encoding if r.encoding and r.encoding.lower() not in ("iso-8859-1", "ascii") else "utf-8"
    return r.url, r.content.decode(enc, errors="replace")


def robots_ok(base: str) -> RobotFileParser | None:
    rp = RobotFileParser()
    try:
        r = session().get(urljoin(base, "/robots.txt"), timeout=10)
        if r.status_code >= 400:
            return None
        rp.parse(r.text.splitlines())
        return rp
    except Exception:
        return None


def page_text(soup: BeautifulSoup) -> str:
    for t in soup(["script", "style", "noscript", "svg", "iframe"]):
        t.decompose()
    return re.sub(r"\s+", " ", soup.get_text(" ")).strip()


# ---------------------------------------------------------------- scoring

@dataclass
class Criterion:
    key: str
    label: str
    weight: float
    saturate: int
    patterns: list


def load_criteria(path):
    cfg = json.load(open(path, encoding="utf-8"))
    crits = [Criterion(k, v["label"], v["weight"], v.get("saturate", 3),
                       [re.compile(p, re.I) for p in v["patterns"]])
             for k, v in cfg["criteria"].items()]
    return crits, cfg.get("fit_thresholds", {"strong": 65, "possible": 40}), cfg.get("require_for_strong", [])


def snippet(text, m, width=110):
    a, b = max(0, m.start() - width), min(len(text), m.end() + width)
    return ("…" if a else "") + text[a:b].strip() + ("…" if b < len(text) else "")


def score_pages(pages, crits, thresholds, required):
    """pages: list of (url, text). Returns dict of results."""
    out, total = {}, 0.0
    for c in crits:
        matched, evidence = [], ""
        for p in c.patterns:
            for url, text in pages:
                m = p.search(text)
                if m:
                    matched.append(m.group(0).strip())
                    if not evidence:
                        evidence = f"{snippet(text, m)}  [{url}]"
                    break
        frac = min(len(matched), c.saturate) / c.saturate
        total += frac * c.weight
        out[c.key] = {"score": round(frac * c.weight, 1), "terms": sorted(set(t.lower() for t in matched)),
                      "evidence": evidence}
    total = round(total, 1)
    if total >= thresholds["strong"] and all(out[k]["score"] > 0 for k in required):
        fit = "Strong"
    elif total >= thresholds["possible"]:
        fit = "Possible"
    else:
        fit = "Weak"
    return total, fit, out


# ---------------------------------------------------------------- crawling

@dataclass
class Org:
    domain: str
    url: str
    name: str = ""
    source: str = "seed"
    depth: int = 0
    linked_by: set = field(default_factory=set)
    title: str = ""
    description: str = ""
    pages: list = field(default_factory=list)
    emails: set = field(default_factory=set)
    phones: set = field(default_factory=set)
    outbound: set = field(default_factory=set)
    error: str = ""
    total: float = 0.0
    fit: str = ""
    results: dict = field(default_factory=dict)


def guess_name(soup, fallback):
    og = soup.find("meta", property="og:site_name")
    if og and og.get("content"):
        return og["content"].strip()
    title = (soup.title.string or "").strip() if soup.title else ""
    if title:
        parts = [p.strip() for p in re.split(r"\s[|\-–—:]\s", title) if p.strip()]
        parts = [p for p in parts if p.lower() not in ("home", "homepage", "welcome")] or parts
        return min(parts, key=len) if len(parts) > 1 and len(min(parts, key=len)) > 3 else parts[0]
    return fallback


def collect_links(soup, base_url, own_domain):
    internal, outbound_all, outbound_partner = [], set(), set()
    page_is_partnerish = bool(PARTNERISH.search(base_url))
    for a in soup.find_all("a", href=True):
        href = a["href"].strip()
        if href.startswith(("mailto:", "tel:", "javascript:", "#")):
            continue
        u = urljoin(base_url, href).split("#")[0]
        if not u.startswith("http") or SKIP_EXT.search(u):
            continue
        d = reg_domain(u)
        if d == own_domain:
            label = (a.get_text(" ") or "") + " " + urlparse(u).path
            if RELEVANT.search(label):
                internal.append((u, bool(PARTNERISH.search(label))))
        elif not is_blocked(d):
            outbound_all.add(d)
            if page_is_partnerish or a.find("img") is not None:  # partner logos are usually images
                outbound_partner.add(d)
    return internal, outbound_all, outbound_partner


def crawl(org: Org, max_pages: int, delay: float):
    try:
        final, html = fetch(org.url)
    except Exception as e:
        # retry with/without www
        alt = org.url.replace("://www.", "://") if "://www." in org.url else org.url.replace("://", "://www.")
        try:
            final, html = fetch(alt)
        except Exception:
            org.error = f"{type(e).__name__}: {str(e)[:120]}"
            return org
    org.url = final
    own = reg_domain(final)
    rp = robots_ok(final)
    soup = BeautifulSoup(html, "html.parser")
    org.title = (soup.title.string or "").strip() if soup.title and soup.title.string else ""
    md = soup.find("meta", attrs={"name": "description"}) or soup.find("meta", property="og:description")
    org.description = (md.get("content") or "").strip() if md else ""
    if not org.name:
        org.name = guess_name(soup, org.domain)
    internal, out_all, out_partner = collect_links(soup, final, own)
    texts = [(final, page_text(soup))]
    raw_html = [html]

    # prioritise partner pages first (for snowballing), then other relevant pages
    seen, queue = {final.rstrip("/")}, []
    for u, partnerish in sorted(internal, key=lambda x: not x[1]):
        k = u.rstrip("/")
        if k not in seen:
            seen.add(k)
            queue.append(u)
    for u in queue[: max_pages - 1]:
        if rp and not rp.can_fetch(UA, u):
            continue
        time.sleep(delay)
        try:
            fu, h = fetch(u)
        except Exception:
            continue
        s = BeautifulSoup(h, "html.parser")
        _, oa, op = collect_links(s, fu, own)
        out_all |= oa
        out_partner |= op
        texts.append((fu, page_text(s)))
        raw_html.append(h)

    org.pages = texts
    blob = " ".join(raw_html)
    org.emails = {e for e in EMAIL_RE.findall(blob)
                  if not re.search(r"\.(png|jpe?g|gif|webp|svg)$|sentry|example\.|wixpress|@\dx", e, re.I)}
    org.phones = {re.sub(r"\s+", " ", p) for p in PHONE_RE.findall(" ".join(t for _, t in texts))}
    org.outbound = out_partner or out_all
    return org


# ---------------------------------------------------------------- search (optional)

def serper_search(queries, key, per_query=10):
    found = {}
    for q in queries:
        try:
            r = requests.post("https://google.serper.dev/search",
                              headers={"X-API-KEY": key, "Content-Type": "application/json"},
                              json={"q": q, "gl": "in", "num": per_query}, timeout=30)
            r.raise_for_status()
            for item in r.json().get("organic", []):
                link = item.get("link")
                if not link:
                    continue
                d = reg_domain(link)
                if not is_blocked(d):
                    found.setdefault(d, (f"https://{urlparse(link).netloc}", q))
        except Exception as e:
            print(f"  search failed for '{q}': {e}")
    return found


def load_search_queries(path):
    """Flatten every tag's optional 'search_queries' list out of params_library.json, deduped."""
    cfg = json.load(open(path, encoding="utf-8"))
    seen, out = set(), []
    for cat in cfg.get("categories", {}).values():
        for tag in cat.get("tags", {}).values():
            for q in tag.get("search_queries", []):
                if q not in seen:
                    seen.add(q)
                    out.append(q)
    return out


# ---------------------------------------------------------------- output

def write_excel(orgs, crits, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Matches"
    ok = sorted([o for o in orgs if not o.error], key=lambda o: -o.total)
    head = ["Rank", "Organisation", "Website", "Fit", "Score /100"] + [c.label for c in crits] + \
           ["Matched terms"] + [f"Evidence: {c.label}" for c in crits] + \
           ["Description", "Emails", "Phones", "Found via", "Linked by", "Pages read", "Note"]
    ws.append(head)
    fills = {"Strong": "C6EFCE", "Possible": "FFEB9C", "Weak": "F2F2F2"}
    for i, o in enumerate(ok, 1):
        terms = "; ".join(f"{c.key}: {', '.join(o.results[c.key]['terms'])}" for c in crits if o.results[c.key]["terms"])
        row = [i, o.name, o.url, o.fit, o.total] + [o.results[c.key]["score"] for c in crits] + [terms] + \
              [o.results[c.key]["evidence"] for c in crits] + \
              [o.description[:400], ", ".join(sorted(o.emails))[:300], ", ".join(sorted(o.phones))[:200],
               o.source, ", ".join(sorted(o.linked_by))[:300], len(o.pages),
               "Little readable text (JavaScript site?) - check manually"
               if sum(len(t) for _, t in o.pages) < 2000 else ""]
        ws.append(row)
        ws.cell(row=i + 1, column=4).fill = PatternFill("solid", fgColor=fills[o.fit])
        ws.cell(row=i + 1, column=3).hyperlink = o.url
    widths = {"Organisation": 32, "Website": 36, "Description": 60, "Matched terms": 50}
    for j, h in enumerate(head, 1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(h, 70 if h.startswith("Evidence") else 14)
        ws.cell(row=1, column=j).font = Font(bold=True)
        ws.cell(row=1, column=j).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    es = wb.create_sheet("Unreachable")
    es.append(["Organisation", "URL", "Found via", "Error"])
    for o in orgs:
        if o.error:
            es.append([o.name, o.url, o.source, o.error])
    for col, w in zip("ABCD", (30, 40, 30, 80)):
        es.column_dimensions[col].width = w
    wb.save(path)


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--seeds", default=os.path.join(HERE, "seeds.csv"))
    ap.add_argument("--criteria", default=os.path.join(HERE, "criteria.json"))
    ap.add_argument("--depth", type=int, default=1, help="snowball levels beyond seeds (0 = seeds only)")
    ap.add_argument("--max-orgs", type=int, default=250, help="cap on total organisations crawled")
    ap.add_argument("--min-links", type=int, default=1,
                    help="only snowball to orgs linked by at least this many crawled orgs")
    ap.add_argument("--max-pages", type=int, default=8, help="pages read per website")
    ap.add_argument("--workers", type=int, default=8)
    ap.add_argument("--delay", type=float, default=0.5, help="seconds between requests to the same site")
    ap.add_argument("--search", action="store_true", help="also discover via Serper (needs SERPER_API_KEY)")
    ap.add_argument("--params", default=os.path.join(HERE, "params_library.json"),
                    help="where --search reads its query phrases from (each tag's 'search_queries')")
    ap.add_argument("--out", default=None)
    ap.add_argument("--db", default=os.path.join(HERE, "orgs.db"),
                    help="SQLite database to upsert results into (shared across runs); pass '' to disable")
    ap.add_argument("--no-excel", action="store_true", help="skip writing the per-run Excel snapshot")
    ap.add_argument("--retry-errors", action="store_true",
                    help="also re-crawl domains already in the database that previously errored "
                         "(by default, anything already in the database is never re-crawled)")
    a = ap.parse_args()

    started = datetime.now(timezone.utc).isoformat(timespec="seconds")
    crits, thresholds, required = load_criteria(a.criteria)
    orgs: dict[str, Org] = {}
    with open(a.seeds, encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            url = row["url"].strip()
            if not url:
                continue
            if not url.startswith("http"):
                url = "https://" + url
            d = reg_domain(url)
            orgs.setdefault(d, Org(domain=d, url=url, name=row.get("name", "").strip(), source="seed"))

    if a.search:
        key = os.environ.get("SERPER_API_KEY")
        if not key:
            sys.exit("--search needs SERPER_API_KEY set (free key at https://serper.dev)")
        queries = load_search_queries(a.params)
        if not queries:
            sys.exit(f"No search_queries found in {a.params} (add some to tags' 'search_queries' lists)")
        print(f"Searching ({len(queries)} queries)…")
        for d, (u, q) in serper_search(queries, key).items():
            orgs.setdefault(d, Org(domain=d, url=u, source=f"search: {q}"))

    db_con = orgdb.open_db(a.db) if a.db else None

    known = orgdb.all_domains(db_con) if db_con is not None else set()
    if a.retry_errors and db_con is not None:
        errored = {r[0] for r in db_con.execute("SELECT domain FROM orgs WHERE error != ''")}
        known -= errored
    if known:
        skip = [d for d in orgs if d in known]
        for d in skip:
            del orgs[d]
        print(f"Skipping {len(skip)} organisation(s) already in the database "
              f"(pass --retry-errors to retry ones that previously failed)")

    crawled: set[str] = set()
    frontier = list(orgs.values())
    for level in range(a.depth + 1):
        budget = a.max_orgs - len(crawled)
        frontier = frontier[:max(budget, 0)]
        if not frontier:
            break
        print(f"\nLevel {level}: crawling {len(frontier)} sites…")
        done = 0
        with ThreadPoolExecutor(max_workers=a.workers) as ex:
            futs = {ex.submit(crawl, o, a.max_pages, a.delay): o for o in frontier}
            for fut in as_completed(futs):
                o = fut.result()
                crawled.add(o.domain)
                done += 1
                if not o.error:
                    o.total, o.fit, o.results = score_pages(o.pages, crits, thresholds, required)
                    print(f"  [{done}/{len(frontier)}] {o.fit:8} {o.total:5.1f}  {o.name[:45]}  ({o.domain})")
                else:
                    print(f"  [{done}/{len(frontier)}] ERROR           {o.domain}: {o.error[:60]}")
                if db_con is not None:
                    orgdb.upsert_org(db_con, {
                        "domain": o.domain, "url": o.url, "name": o.name, "title": o.title,
                        "description": o.description, "emails": o.emails, "phones": o.phones,
                        "pages": o.pages, "source": o.source, "linked_by": o.linked_by, "error": o.error,
                        "now": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                    })

        if level == a.depth:
            break
        # snowball: candidates linked from orgs that scored at least "possible"
        linkers = [o for o in frontier if not o.error and o.total >= thresholds["possible"]]
        cand: dict[str, Org] = {}
        for o in linkers:
            for d in o.outbound:
                if d in crawled or d in known or d in orgs and orgs[d].domain in crawled:
                    continue
                c = cand.setdefault(d, orgs.get(d) or Org(domain=d, url=f"https://{d}", depth=level + 1))
                c.linked_by.add(o.name or o.domain)
        for d, c in cand.items():
            c.source = f"linked from {len(c.linked_by)} org(s)"
            orgs[d] = c
        frontier = sorted((c for c in cand.values() if len(c.linked_by) >= a.min_links),
                          key=lambda c: -len(c.linked_by))

    results = [o for o in orgs.values() if o.domain in crawled]
    n = {f: sum(1 for o in results if o.fit == f and not o.error) for f in ("Strong", "Possible", "Weak")}
    n_error = sum(1 for o in results if o.error)
    print(f"\nDone. {len(results)} sites crawled — Strong: {n['Strong']}, Possible: {n['Possible']}, "
          f"Weak: {n['Weak']}, Unreachable: {n_error}")

    if db_con is not None:
        orgdb.log_run(db_con, started, datetime.now(timezone.utc).isoformat(timespec="seconds"),
                      vars(a), len(results), n_error)
        total_in_db = len(orgdb.all_domains(db_con))
        db_con.commit()
        db_con.close()
        print(f"Saved to database: {a.db}  (now {total_in_db} organisations total — "
              f"run the query app to filter/rank by parameters)")

    if not a.no_excel:
        out = a.out or os.path.join(HERE, f"results_{datetime.now():%Y%m%d_%H%M}.xlsx")
        write_excel(results, crits, out)
        print(f"Saved snapshot: {out}")


if __name__ == "__main__":
    main()
